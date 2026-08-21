"""excel_exporter.py — Generate multi-sheet Excel report for linter revision benchmark."""

import json
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HDR_FONT = Font(bold=True, size=11, color="FFFFFF")
GOOD = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BAD = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
SEP_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
CTR = Alignment(horizontal="center")
LEFT = Alignment(horizontal="left", wrap_text=True)


def _cell(ws, row, col, val, fill=None, align=CTR):
    c = ws.cell(row=row, column=col, value=val)
    c.alignment = align
    c.border = BORDER
    if fill:
        c.fill = fill
    return c


def _header(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CTR
        c.border = BORDER


def _pct(val):
    if val is None:
        return "N/A"
    return f"{val * 100:.1f}%"


def _f2(val):
    if val is None:
        return "N/A"
    return f"{val:.2f}"


def _auto_width(ws, min_width=10, max_width=50):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def build_readme_config(ws, config: Dict):
    headers = ["Parameter", "Value"]
    _header(ws, 1, headers)
    rows = [
        ("Raw data path", config.get("raw_data_path", "")),
        ("Linter tool", "clang-tidy"),
        ("Clang-tidy checks", ", ".join(config.get("checks", []))),
        ("C++ standard", config.get("cpp_standard", "c++17")),
        ("Train/Valid/Test ratio", f"{config.get('train_ratio', 0.7)}/{config.get('valid_ratio', 0.15)}/{config.get('test_ratio', 0.15)}"),
        ("Random seed", str(config.get("random_seed", 42))),
        ("Max submissions per problem", str(config.get("max_submissions_per_problem", 1))),
        ("DeepSeek model", config.get("deepseek_model", "deepseek-v3.2-128k")),
        ("Temperature", str(config.get("temperature", 0.2))),
        ("Min LOC filter", str(config.get("min_loc", 5))),
        ("Max LOC filter", str(config.get("max_loc", 500))),
        ("Min messages filter", str(config.get("min_messages", 1))),
        ("Max messages filter", str(config.get("max_messages", 200))),
        ("Max line churn filter", str(config.get("max_line_churn", 0.8))),
        ("Max new messages filter", str(config.get("max_new_messages", 10))),
    ]
    for i, (param, val) in enumerate(rows, 2):
        _cell(ws, i, 1, param, align=LEFT)
        _cell(ws, i, 2, val)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_problems_sheet(ws, problems: List[Dict], split_map: Dict):
    headers = ["Problem ID", "Source", "Title", "Time Limit (ms)", "Memory Limit (MB)",
               "Num Correct Subs", "Num C++ Subs", "Split"]
    _header(ws, 1, headers)
    for i, p in enumerate(problems, 2):
        cpp_count = sum(1 for s in p.get("correct_submissions", []) if s.get("language") == "cpp")
        _cell(ws, i, 1, p.get("id", ""), align=LEFT)
        _cell(ws, i, 2, p.get("source", ""))
        _cell(ws, i, 3, p.get("title", ""), align=LEFT)
        _cell(ws, i, 4, p.get("time_limit", ""))
        _cell(ws, i, 5, p.get("memory_limit", ""))
        _cell(ws, i, 6, len(p.get("correct_submissions", [])))
        _cell(ws, i, 7, cpp_count)
        _cell(ws, i, 8, split_map.get(p.get("id", ""), "unknown"))
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_candidates_sheet(ws, candidates: List[Dict]):
    headers = ["Sample ID", "Problem ID", "Sub Index", "Language", "Original LOC",
               "Compile OK", "Test Gen OK", "Original Total Messages",
               "Filter Status", "Filter Reason"]
    _header(ws, 1, headers)
    for i, c in enumerate(candidates, 2):
        _cell(ws, i, 1, c.get("sample_id", ""))
        _cell(ws, i, 2, c.get("problem_id", ""), align=LEFT)
        _cell(ws, i, 3, c.get("submission_index", ""))
        _cell(ws, i, 4, c.get("language", ""))
        _cell(ws, i, 5, c.get("original_loc", ""))
        compile_ok = c.get("compile_ok")
        _cell(ws, i, 6, str(compile_ok) if compile_ok is not None else "",
              GOOD if compile_ok else BAD if compile_ok is False else None)
        test_ok = c.get("test_generation_ok")
        _cell(ws, i, 7, str(test_ok) if test_ok is not None else "",
              GOOD if test_ok else BAD if test_ok is False else None)
        _cell(ws, i, 8, c.get("original_total_messages", ""))
        status = c.get("filter_status", "")
        _cell(ws, i, 9, status, GOOD if status == "accepted" else BAD if status == "rejected" else None)
        _cell(ws, i, 10, c.get("filter_reason", ""), align=LEFT)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_linter_before_sheet(ws, all_linter_messages: Dict):
    headers = ["Sample ID", "Msg Index", "Tool", "Category", "Check Name",
               "Severity", "Line", "Column", "Message Text", "Message Identity"]
    _header(ws, 1, headers)
    row = 2
    for sample_id, messages in all_linter_messages.items():
        for idx, m in enumerate(messages):
            _cell(ws, row, 1, sample_id)
            _cell(ws, row, 2, idx)
            _cell(ws, row, 3, m.get("tool", ""))
            _cell(ws, row, 4, m.get("category", ""))
            _cell(ws, row, 5, m.get("check_name", ""), align=LEFT)
            _cell(ws, row, 6, m.get("severity", ""))
            _cell(ws, row, 7, m.get("line", ""))
            _cell(ws, row, 8, m.get("column", ""))
            _cell(ws, row, 9, m.get("message_text", ""), align=LEFT)
            _cell(ws, row, 10, m.get("message_identity", ""), align=LEFT)
            row += 1
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_revisions_sheet(ws, revisions: List[Dict]):
    headers = ["Sample ID", "DeepSeek Model", "Prompt Hash",
               "Extraction OK", "Generation Error"]
    _header(ws, 1, headers)
    for i, r in enumerate(revisions, 2):
        _cell(ws, i, 1, r.get("sample_id", ""))
        _cell(ws, i, 2, r.get("deepseek_model", ""))
        _cell(ws, i, 3, r.get("prompt_hash", ""))
        ext_ok = r.get("extraction_ok")
        _cell(ws, i, 4, str(ext_ok) if ext_ok is not None else "",
              GOOD if ext_ok else BAD if ext_ok is False else None)
        _cell(ws, i, 5, r.get("generation_error", ""), align=LEFT)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_validation_sheet(ws, validations: List[Dict]):
    headers = ["Sample ID", "Compile OK", "Pass@1", "Num Tests",
               "Tests Passed", "Tests Failed", "Checker Available",
               "Checker Status", "Runtime Error", "Timeout"]
    _header(ws, 1, headers)
    for i, v in enumerate(validations, 2):
        _cell(ws, i, 1, v.get("sample_id", ""))
        compile_ok = v.get("compile_ok")
        _cell(ws, i, 2, str(compile_ok) if compile_ok is not None else "",
              GOOD if compile_ok else BAD if compile_ok is False else None)
        pass1 = v.get("pass_at_1")
        _cell(ws, i, 3, str(pass1) if pass1 is not None else "",
              GOOD if pass1 else BAD if pass1 is False else None)
        _cell(ws, i, 4, v.get("tests_generated", 0))
        _cell(ws, i, 5, v.get("tests_passed", 0))
        _cell(ws, i, 6, v.get("tests_failed", 0))
        _cell(ws, i, 7, str(v.get("checker_available", "")))
        _cell(ws, i, 8, v.get("checker_status", ""))
        _cell(ws, i, 9, str(v.get("runtime_error", "")), align=LEFT)
        _cell(ws, i, 10, str(v.get("timeout", "")))
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_metrics_sheet(ws, metrics: List[Dict]):
    headers = ["Sample ID", "Original Total Msgs", "Revised Total Msgs",
               "Old Msgs Remaining", "New Msgs Count",
               "MRR@Pass ↑", "New Msg./Sample@Pass ↓", "Line Churn@Pass ↓",
               "Original LOC", "Revised LOC", "Added Lines", "Deleted Lines"]
    _header(ws, 1, headers)
    for i, m in enumerate(metrics, 2):
        _cell(ws, i, 1, m.get("sample_id", ""))
        _cell(ws, i, 2, m.get("original_total_messages", 0))
        _cell(ws, i, 3, m.get("revised_total_messages", 0))
        _cell(ws, i, 4, m.get("old_messages_remaining", 0))
        _cell(ws, i, 5, m.get("new_messages_count", 0))
        mrr = m.get("mrr_at_pass")
        _cell(ws, i, 6, _pct(mrr),
              GOOD if mrr is not None and mrr > 0 else BAD if mrr is not None and mrr < -0.3 else None)
        new_msg = m.get("new_msg_per_sample_at_pass")
        _cell(ws, i, 7, _f2(new_msg),
              GOOD if new_msg is not None and new_msg < 1 else None)
        churn = m.get("line_churn_at_pass")
        _cell(ws, i, 8, _f2(churn))
        _cell(ws, i, 9, m.get("original_loc", 0))
        _cell(ws, i, 10, m.get("revised_loc", 0))
        _cell(ws, i, 11, m.get("added_lines", 0))
        _cell(ws, i, 12, m.get("deleted_lines", 0))
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_train_pairs_sheet(ws, pairs: List[Dict]):
    headers = ["Sample ID", "Problem ID", "Split", "Include in Train",
               "Include Reason", "Input Preview", "Output Preview"]
    _header(ws, 1, headers)
    for i, p in enumerate(pairs, 2):
        _cell(ws, i, 1, p.get("sample_id", ""))
        _cell(ws, i, 2, p.get("problem_id", ""), align=LEFT)
        _cell(ws, i, 3, p.get("split", ""))
        include = p.get("include_in_train", False)
        _cell(ws, i, 4, str(include), GOOD if include else BAD)
        _cell(ws, i, 5, p.get("include_reason", ""), align=LEFT)
        inp = p.get("input_preview", "")
        _cell(ws, i, 6, inp[:200] if inp else "", align=LEFT)
        out = p.get("output_preview", "")
        _cell(ws, i, 7, out[:200] if out else "", align=LEFT)
    ws.freeze_panes = "A2"
    _auto_width(ws)


def build_summary_sheet(ws, summary: Dict):
    headers = ["Split", "Num Candidates", "Num Valid Pairs",
               "Pass@1", "MRR@Pass ↑", "New Msg./Sample@Pass ↓",
               "Line Churn@Pass ↓", "Avg Original Msgs",
               "Bugprone MRR", "Performance MRR", "Modernize MRR", "Readability MRR", "Other MRR"]
    _header(ws, 1, headers)
    for i, (split_name, s) in enumerate(summary.items(), 2):
        _cell(ws, i, 1, split_name, align=LEFT)
        _cell(ws, i, 2, s.get("num_candidates", 0))
        _cell(ws, i, 3, s.get("num_valid_pairs", 0))
        _cell(ws, i, 4, _pct(s.get("pass_rate")))
        _cell(ws, i, 5, _pct(s.get("avg_mrr_at_pass")))
        _cell(ws, i, 6, _f2(s.get("avg_new_msg_at_pass")))
        _cell(ws, i, 7, _f2(s.get("avg_line_churn_at_pass")))
        _cell(ws, i, 8, _f2(s.get("avg_original_messages")))
        cat_mrr = s.get("category_mrr", {})
        for col, cat in enumerate(["bugprone", "performance", "modernize", "readability", "other"], 9):
            v = cat_mrr.get(cat)
            _cell(ws, i, col, _pct(v))
    ws.freeze_panes = "A2"
    _auto_width(ws)


def export_excel(output_path: str, config: Dict, problems: List[Dict],
                 candidates: List[Dict], split_map: Dict,
                 linter_messages: Dict, revisions: List[Dict],
                 validations: List[Dict], metrics: List[Dict],
                 train_pairs: List[Dict], summary: Dict):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "README_Config"
    build_readme_config(ws1, config)

    ws2 = wb.create_sheet("Problems")
    build_problems_sheet(ws2, problems, split_map)

    ws3 = wb.create_sheet("CandidateSubmissions")
    build_candidates_sheet(ws3, candidates)

    ws4 = wb.create_sheet("LinterBefore")
    build_linter_before_sheet(ws4, linter_messages)

    ws5 = wb.create_sheet("GeneratedRevisions")
    build_revisions_sheet(ws5, revisions)

    ws6 = wb.create_sheet("ValidationResults")
    build_validation_sheet(ws6, validations)

    ws7 = wb.create_sheet("Metrics")
    build_metrics_sheet(ws7, metrics)

    ws8 = wb.create_sheet("FinalTrainPairs")
    build_train_pairs_sheet(ws8, train_pairs)

    ws9 = wb.create_sheet("Summary")
    build_summary_sheet(ws9, summary)

    wb.save(output_path)
    print(f"Excel saved: {output_path} ({len(wb.sheetnames)} sheets)")
